from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def create_presentation():
    prs = Presentation()

    def add_bullet_slide(title, bullet_points):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        for point in bullet_points:
            p = tf.add_paragraph()
            p.text = point
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(0, 0, 0)
            p.alignment = PP_ALIGN.LEFT

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Защита проекта: Django приложение"
    subtitle = slide.placeholders[1]
    subtitle.text = "Автор: Ваше имя\nДата: 2024"

    add_bullet_slide("Введение", [
        "Описание проекта: интернет-магазин спортивных товаров",
        "Цель: создать удобный и функциональный веб-сайт для продажи товаров"
    ])

    add_bullet_slide("Актуальность и мотивация", [
        "Рост онлайн-торговли",
        "Потребность в удобных решениях для спорта"
    ])

    add_bullet_slide("Цели и задачи проекта", [
        "Разработка модели данных",
        "Создание пользовательского интерфейса",
        "Реализация корзины и оплаты"
    ])

    add_bullet_slide("Архитектура проекта", [
        "Использован Django 5.2",
        "REST API с DRF",
        "PostgreSQL в качестве базы данных"
    ])

    add_bullet_slide("Подключение к базе данных", [
        "Используется PostgreSQL",
        "Настройки в nxt_sports/settings.py:",
        "ENGINE: django.db.backends.postgresql",
        "NAME: nxt_sports_db",
        "USER: postgres",
        "HOST: localhost",
        "PORT: 5432"
    ])

    add_bullet_slide("Модели данных", [
        "UserProfile: расширение стандартного пользователя",
        "ProductCategory и Product: категории и товары",
        "Cart и CartItem: корзина и позиции в ней"
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Диаграмма моделей Django"
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(8)
    height = Inches(4.5)
    png_path = os.path.join(os.getcwd(), "my_project_models.png")
    if os.path.exists(png_path):
        slide.shapes.add_picture(png_path, left, top, width=width, height=height)
    else:
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.text = "Диаграмма моделей не найдена.\nПожалуйста, положите my_project_models.png в корень проекта."

    add_bullet_slide("Основные функции и возможности", [
        "Просмотр каталога товаров",
        "Управление корзиной",
        "Профиль пользователя с платежными данными"
    ])

    add_bullet_slide("Интерфейс пользователя", [
        "Адаптивный дизайн",
        "Простая навигация",
        "Страницы: каталог, корзина, профиль"
    ])

    add_bullet_slide("API и сериализация", [
        "Использован Django REST Framework",
        "Сериализаторы для моделей",
        "REST API для взаимодействия с фронтендом"
    ])

    add_bullet_slide("Тестирование", [
        "Юнит-тесты для моделей и представлений",
        "Тесты API",
        "Использование pytest"
    ])

    add_bullet_slide("Безопасность и аутентификация", [
        "Расширение стандартного пользователя UserProfile",
        "Аутентификация через Django",
        "Защита данных пользователя"
    ])

    add_bullet_slide("Развертывание и настройка", [
        "Использование PostgreSQL",
        "Настройка окружения в settings.py",
        "Миграции и управление базой"
    ])

    add_bullet_slide("Проблемы и решения", [
        "Настройка платежных данных пользователя",
        "Оптимизация запросов к базе",
        "Обработка ошибок и валидация"
    ])

    add_bullet_slide("Результаты и достижения", [
        "Рабочий интернет-магазин",
        "Полное покрытие тестами",
        "Удобный интерфейс"
    ])

    add_bullet_slide("Планы на будущее", [
        "Добавить оплату онлайн",
        "Расширить каталог товаров",
        "Улучшить UX/UI"
    ])

    add_bullet_slide("Заключение", [
        "Проект успешно реализован",
        "Готов к дальнейшему развитию"
    ])

    add_bullet_slide("Вопросы и ответы", [
        "Готов ответить на ваши вопросы"
    ])

    prs.save("project_presentation_filled.pptx")

if __name__ == "__main__":
    create_presentation()
